from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models.leads import Lead, LeadStatus
from app.services.excel_service import generate_leads_excel


HEADERS = [
    "Name",
    "Email",
    "Contact Number",
    "Purpose of Contact",
    "Status",
    "Created Date",
]


def _load_exported_worksheet(lead: Lead):
    output = generate_leads_excel([lead])
    workbook = load_workbook(BytesIO(output.getvalue()), data_only=False)
    return workbook["Leads"]


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("=1+1", "=1+1"),
        ("+1+1", "+1+1"),
        ("-1+1", "-1+1"),
        ("@SUM(1,1)", "@SUM(1,1)"),
        (" =1+1", " =1+1"),
        ("\t=1+1", "\t=1+1"),
        ("\n=1+1", "\n=1+1"),
        ("\r=1+1", "\n=1+1"),
        (" \t=1+1", " \t=1+1"),
    ],
)
def test_formula_like_lead_names_are_written_as_text(value, expected):
    worksheet = _load_exported_worksheet(
        Lead(
            name=value,
            email="safe@example.com",
            phone="0501234567",
            service="Home Financing",
            message="Synthetic test lead",
            status=LeadStatus.NEW,
            created_at=datetime(2025, 1, 2, 3, 4),
        )
    )

    cell = worksheet["A2"]
    # XML normalizes carriage returns to line feeds when the workbook is read.
    assert cell.value == expected
    assert cell.data_type == "s"


def test_all_public_text_fields_are_written_as_strings():
    worksheet = _load_exported_worksheet(
        Lead(
            name="=1+1",
            email="+formula@example.com",
            phone="-1234567",
            service="@SUM(1,1)",
            message="Synthetic test lead",
            status=LeadStatus.NEW,
            created_at=datetime(2025, 1, 2, 3, 4),
        )
    )

    for coordinate, expected in {
        "A2": "=1+1",
        "B2": "+formula@example.com",
        "C2": "-1234567",
        "D2": "@SUM(1,1)",
    }.items():
        cell = worksheet[coordinate]
        assert cell.value == expected
        assert cell.data_type == "s"


def test_workbook_structure_and_legitimate_values_are_preserved():
    created_at = datetime(2025, 1, 2, 3, 4)
    worksheet = _load_exported_worksheet(
        Lead(
            name="Alex Example",
            email="alex@example.com",
            phone="0501234567",
            service="Home Financing",
            message="Synthetic test lead",
            status=LeadStatus.NEW,
            created_at=created_at,
        )
    )

    assert [cell.value for cell in worksheet[1]] == HEADERS
    assert [worksheet.cell(row=2, column=index).value for index in range(1, 7)] == [
        "Alex Example",
        "alex@example.com",
        "0501234567",
        "Home Financing",
        "NEW",
        created_at,
    ]
    assert worksheet["C2"].data_type == "s"
    assert worksheet["F2"].data_type == "d"
    assert worksheet["F2"].number_format == "DD-MMM-YYYY HH:MM"
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == "A1:F2"


def test_empty_text_fields_remain_empty_and_non_formula():
    worksheet = _load_exported_worksheet(
        Lead(
            name="",
            email="",
            phone="",
            service="",
            message="Synthetic test lead",
            status=LeadStatus.NEW,
            created_at=datetime(2025, 1, 2, 3, 4),
        )
    )

    for coordinate in ("A2", "B2", "C2", "D2"):
        cell = worksheet[coordinate]
        assert cell.value in ("", None)
        assert cell.data_type != "f"
