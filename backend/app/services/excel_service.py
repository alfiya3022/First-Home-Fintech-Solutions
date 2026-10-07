from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.leads import Lead


def generate_leads_excel(leads: list[Lead]) -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Leads"

    headers = [
        "Name",
        "Email",
        "Contact Number",
        "Purpose of Contact",
        "Status",
        "Created Date",
    ]

    worksheet.append(headers)

    # Header styling
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="071B2D",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    # Lead data
    for lead in leads:
        worksheet.append(
            [
                lead.name,
                lead.email,
                lead.phone,
                lead.service,
                lead.status.value,
                lead.created_at,
            ]
        )

    # Date formatting
    for row in worksheet.iter_rows(
        min_row=2,
        min_col=6,
        max_col=6,
    ):
        row[0].number_format = "DD-MMM-YYYY HH:MM"

    # Column widths
    column_widths = {
        1: 24,
        2: 32,
        3: 20,
        4: 38,
        5: 16,
        6: 24,
    }

    for column_index, width in column_widths.items():
        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = width

    # Freeze header row
    worksheet.freeze_panes = "A2"

    # Enable filtering
    worksheet.auto_filter.ref = worksheet.dimensions

    output = BytesIO()

    workbook.save(output)
    output.seek(0)

    return output